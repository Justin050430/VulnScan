# -*- coding: utf-8 -*-
"""Inspect the CURRENT FYP.docx (authoritative, _extracted text files are stale)."""
import docx
from PIL import Image

PATH = r"c:\FYP\VulnScan\Assignment\FYP.docx"
d = docx.Document(PATH)

print("=" * 90)
print("PARAGRAPH COUNT:", len(d.paragraphs), " TABLE COUNT:", len(d.tables))
print("=" * 90)

KEYWORDS = (
    "Deployment Architecture", "Architecture Design", "Database Design",
    "Figure 4.1:", "Figure 4.2:", "Figure 4.3:", "Figure 4.4:",
    "Non-Functional Requirements of VulnScan", "Compatibility", "Availability",
    "Microsoft SQL Server was selected", "NFR-07", "NFR-08", "NFR-09",
    "system is designed as a client-server", "Azure", "deployment",
)

for i, p in enumerate(d.paragraphs):
    t = p.text
    if any(k.lower() in t.lower() for k in KEYWORDS):
        print(f"[{i}][{p.style.name}] {t}")
        print("-" * 90)

print("=" * 90)
print("TABLE 4 (expect NFR Table 3.2):")
for ri, row in enumerate(d.tables[4].rows):
    cells = [c.text.strip() for c in row.cells]
    print(f"  row {ri}: {cells}")
print("=" * 90)
print("TABLE 3 (expect FR Table 3.1) - last 3 rows only:")
for ri, row in enumerate(d.tables[3].rows):
    if ri >= len(d.tables[3].rows) - 3:
        print(f"  row {ri}: {[c.text.strip() for c in row.cells]}")
print("=" * 90)

# ---- figure PNG colour analysis (flat monochrome vs 3D shaded) ----
im = Image.open(r"c:\FYP\VulnScan\docs\design\Fig4.2_Deployment_Diagram.png").convert("RGB")
colors = im.getcolors(maxcolors=1000000)
colors = sorted(colors, reverse=True)
print("Fig4.2 PNG size:", im.size)
print("Distinct colours:", len(colors))
for cnt, rgb in colors[:8]:
    print(f"   {cnt:>8} x {rgb}")
