"""
Align 'OoiJianXin_RIS_Form2.docx' (Project Proposal) with the delivered
VulnScan implementation, keeping the forward-looking proposal voice.

Text-only edits: existing paragraphs / table cells are rewritten in place and
one table row + one reference entry are cloned from existing ones, so all
styles, borders and shading are preserved.
"""
import copy
import shutil

from docx import Document
from docx.oxml.ns import qn
from docx.table import _Row
from docx.text.paragraph import Paragraph

SRC = r"c:\FYP\VulnScan\Assignment\OoiJianXin_RIS_Form2.docx"
BAK = r"c:\FYP\VulnScan\Assignment\OoiJianXin_RIS_Form2_backup_before_edits.docx"

shutil.copyfile(SRC, BAK)
print("backup written:", BAK)

doc = Document(SRC)


def norm(text):
    return " ".join(text.split())


def set_paragraph_text(paragraph, new_text):
    """Rewrite a paragraph keeping its first run's formatting."""
    runs = paragraph.runs
    if not runs:  # pragma: no cover - all targets are plain runs
        paragraph.add_run(new_text)
        return
    runs[0].text = new_text
    for run in runs[1:]:
        run.text = ""


def edit_paragraph(paragraphs, old, new, stats):
    old_n = norm(old)
    for p in paragraphs:
        if norm(p.text) == old_n:
            set_paragraph_text(p, new)
            stats.append(old[:60])
            return True
    stats.append(f"!! NOT FOUND: {old[:60]}")
    return False


STATS = []

# ---------------------------------------------------------------- T1 (main)
t1 = doc.tables[1]
t1_paras = [p for row in t1.rows for c in row.cells for p in c.paragraphs]

edit_paragraph(t1_paras,
    "VulnScan: Development of an Intelligent Web Application Vulnerability Assessment Platform with Risk Assessment and Remediation Knowledge Base",
    "VulnScan \u2013 Intelligent Web Application Vulnerability Assessment "
    "Platform with Risk Assessment and Remediation Knowledge Base",
    STATS)

edit_paragraph(t1_paras,
    "This project proposes the development of VulnScan, an intelligent web application vulnerability assessment system with an integrated security dashboard. The system adopts a modular architecture consisting of a scanning engine, detection engine, risk assessment module and remediation knowledge base. It performs automated crawling, payload injection, and response analysis to detect vulnerabilities, perform vulnerability assessment based on severity and potential impact, and provide remediation guidance through an integrated knowledge base.",
    "This project proposes the development of VulnScan, an intelligent web "
    "application vulnerability assessment system with an integrated security "
    "dashboard. The system adopts a modular architecture consisting of a "
    "scanning engine, detection engine, risk assessment module and remediation "
    "knowledge base. It performs automated crawling, payload injection and "
    "response analysis to detect vulnerabilities; assesses the detected "
    "vulnerabilities according to the Common Vulnerability Scoring System "
    "(CVSS) and their potential impact; and provides remediation guidance "
    "through an integrated knowledge base, supported by configurable scan "
    "profiles, user management and security audit logging.",
    STATS)

edit_paragraph(t1_paras,
    "The proposed system, VulnScan, is a lightweight automated web application vulnerability scanner designed to identify common security weaknesses in web applications within a controlled environment. The system focuses on detecting vulnerabilities such as SQL Injection, Cross-Site Scripting (XSS), and basic authentication flaws through automated scanning and analysis. VulnScan aims to provide a practical and technically robust solution for developers to perform early-stage security assessment without requiring advanced penetration testing expertise. This approach directly addresses the lack of accessible security testing tools, reduces system complexity for developers, and enables earlier detection of vulnerabilities within the development lifecycle.",
    "The proposed system, VulnScan, is a lightweight automated web application "
    "vulnerability scanner designed to identify common security weaknesses in "
    "web applications within a controlled environment. The system focuses on "
    "detecting selected high-priority vulnerabilities, namely SQL Injection "
    "(SQLi), Cross-Site Scripting (XSS) and HTTP security header "
    "misconfigurations, through automated crawling, payload injection and "
    "response analysis. VulnScan aims to provide a practical and technically "
    "robust solution for developers to perform early-stage security assessment "
    "without requiring advanced penetration testing expertise. In addition to "
    "vulnerability detection, the system provides configurable scan profiles, a "
    "vulnerability knowledge base, role-based user management with optional "
    "Google single sign-on, security audit logging and report export, allowing "
    "assessments to be managed and reviewed within a single platform. This "
    "approach directly addresses the lack of accessible security testing tools, "
    "reduces system complexity for developers, and enables earlier detection of "
    "vulnerabilities within the development lifecycle.",
    STATS)

edit_paragraph(t1_paras,
    "Responsible for crawling target web applications and identifying input points such as forms, URLs, and parameters. The engine systematically injects predefined payloads to simulate attack scenarios.",
    "Responsible for crawling target web applications and identifying "
    "accessible pages and input points, such as URL query parameters. The "
    "engine systematically injects predefined payloads into these input points "
    "to simulate attack scenarios, with crawling restricted to the target's own "
    "scope so that no third-party systems are contacted. Support for additional "
    "input types, such as HTML form submission, is identified as a future "
    "enhancement.",
    STATS)

edit_paragraph(t1_paras,
    "Analyses server responses to identify indicators of vulnerabilities. Detection techniques include response pattern analysis, error message identification, and behaviour comparison to determine the presence of SQL Injection and XSS vulnerabilities.",
    "Analyses server responses to identify indicators of vulnerabilities. "
    "Detection techniques include response pattern analysis, error message "
    "identification and behaviour comparison to determine the presence of SQL "
    "Injection, Cross-Site Scripting (XSS) and HTTP security header weaknesses.",
    STATS)

edit_paragraph(t1_paras,
    "Generates structured vulnerability reports including vulnerability type, severity level (Low, Medium, High), affected components, and recommended mitigation strategies.",
    "Generates structured vulnerability reports including the vulnerability "
    "type, severity level and risk score, affected components, supporting "
    "evidence and recommended mitigation strategies. Reports can be exported in "
    "portable formats such as PDF and CSV for documentation and reporting "
    "purposes.",
    STATS)

edit_paragraph(t1_paras,
    "Provides a user interface for managing scans, viewing results, and visualizing detected vulnerabilities through charts and logs. The dashboard also supports scan history tracking and result comparison.",
    "Provides a user interface for managing scans, viewing results and "
    "visualizing detected vulnerabilities through charts and summary "
    "statistics. The dashboard also supports scan history tracking, scan "
    "profile selection and role-based access for Developers and Administrators, "
    "including review of security audit logs.",
    STATS)

edit_paragraph(t1_paras,
    "Assigns a risk score to each detected vulnerability based on severity, frequency, and potential impact. This enables prioritization of critical vulnerabilities and supports more effective decision-making during remediation.",
    "Assigns a severity level and a risk score to each detected vulnerability "
    "using the Common Vulnerability Scoring System (CVSS) v3.1 base metrics "
    "(FIRST, 2019), which take into account the exploitability of the weakness "
    "and its potential impact on confidentiality, integrity and availability. "
    "This enables prioritization of critical vulnerabilities and supports more "
    "effective decision-making during remediation.",
    STATS)

edit_paragraph(t1_paras,
    "Provides actionable mitigation suggestions for each detected vulnerability. For example, in the case of SQL Injection, the system recommends the use of parameterized queries and input validation techniques. This enhances the usability of the system by guiding developers on how to fix identified issues.",
    "Provides actionable mitigation suggestions for each detected "
    "vulnerability. For example, in the case of SQL Injection, the system "
    "recommends the use of parameterized queries and input validation "
    "techniques. Each knowledge base entry documents the vulnerability "
    "description, security impact, remediation recommendation, secure coding "
    "practice and OWASP reference, and administrators can maintain these "
    "entries through the system. This enhances the usability of the system by "
    "guiding developers on how to fix identified issues.",
    STATS)

edit_paragraph(t1_paras,
    "Detection of SQL Injection and XSS vulnerabilities",
    "Detection of SQL Injection, Cross-Site Scripting (XSS) and HTTP security "
    "header weaknesses",
    STATS)

edit_paragraph(t1_paras,
    "Dashboard-based visualization and scan management",
    "Dashboard-based visualization, scan management and security audit log "
    "review",
    STATS)

# ------------------------------------------------- Technology stack (nested)
tech = t1.rows[7].cells[0].tables[0]


def set_cell(cell, old, new):
    edit_paragraph(cell.paragraphs, old, new, STATS)


set_cell(tech.rows[3].cells[1], "MySQL",
         "Microsoft SQL Server (Azure SQL Database)")
set_cell(tech.rows[3].cells[2],
         "Stores scan results, detected vulnerabilities, and historical records for tracking and analysis",
         "Stores user accounts, scan configurations, vulnerability definitions, "
         "scan results and security audit records for tracking and analysis")
set_cell(tech.rows[5].cells[1], "BeautifulSoup",
         "Python Standard Library (html.parser)")
set_cell(tech.rows[5].cells[2],
         "Extracts input fields, forms, and parameters from web pages",
         "Parses HTML responses to extract links, pages and input parameters "
         "during crawling")
set_cell(tech.rows[6].cells[1], "Regular Expressions (Regex)",
         "Response Pattern and Signature Analysis")
set_cell(tech.rows[6].cells[2],
         "Identifies vulnerability patterns based on server responses",
         "Identifies vulnerability indicators through known SQL error "
         "signatures, reflected payload checks, behavioural response comparison "
         "and security header validation")
set_cell(tech.rows[7].cells[1], "Python (Custom Logic)",
         "Python (CVSS v3.1 Base Score)")
set_cell(tech.rows[7].cells[2],
         "Implements risk scoring algorithms and generates remediation recommendations based on detected vulnerabilities",
         "Calculates a standardised risk score and severity level from CVSS "
         "v3.1 base metrics and generates remediation recommendations from the "
         "vulnerability knowledge base")

# add a hosting row by cloning the last row (preserves formatting)
new_tr = copy.deepcopy(tech.rows[-1]._tr)
tech._tbl.append(new_tr)
new_row = _Row(new_tr, tech)
for cell, value in zip(new_row.cells,
                       ("Cloud Deployment",
                        "Microsoft Azure (App Service, Azure SQL Database)",
                        "Hosts the web application and its managed database, with "
                        "automated deployment through a CI/CD pipeline")):
    set_paragraph_text(cell.paragraphs[0], value)
STATS.append("added tech-stack row: Cloud Deployment")

# ------------------------------------------------- Milestones (nested)
ms = doc.tables[2].rows[5].cells[0].tables[0]

edit_paragraph(ms.rows[4].cells[1].paragraphs,
    "Implementation of vulnerability detection logic (SQL Injection, XSS), response analysis module",
    "Implementation of vulnerability detection logic (SQL Injection, XSS and "
    "HTTP security header analysis), risk scoring and response analysis modules",
    STATS)

edit_paragraph(ms.rows[5].cells[1].paragraphs,
    "Integrated backend system using Python (Flask), API endpoints for communication between modules",
    "Integrated backend system using Python (Flask) covering authentication, "
    "user management, scan profiles, the vulnerability knowledge base, security "
    "audit logging and API endpoints for communication between modules",
    STATS)

edit_paragraph(ms.rows[8].cells[1].paragraphs,
    "Complete FYP report (Chapter 1\u20137), user manual, demo preparation and presentation materials",
    "Complete FYP report (Chapter 1\u20137), system deployment, user manual, "
    "demo preparation and presentation materials",
    STATS)

# ------------------------------------------------- in-text citation fix
edit_paragraph(t1_paras,
    "In many development environments, vulnerabilities are only discovered "
    "after deployment or during penetration testing phases, rather than being "
    "identified early in the development lifecycle. This delayed detection "
    "increases the cost and complexity of fixing security issues, as "
    "vulnerabilities may already be embedded deeply within the system. "
    "Furthermore, the absence of automated and simplified scanning solutions "
    "limits the ability to continuously monitor and detect vulnerabilities in a "
    "timely manner, exposing systems to prolonged security risks (NIST, 2007).",
    "In many development environments, vulnerabilities are only discovered "
    "after deployment or during penetration testing phases, rather than being "
    "identified early in the development lifecycle. This delayed detection "
    "increases the cost and complexity of fixing security issues, as "
    "vulnerabilities may already be embedded deeply within the system. "
    "Furthermore, the absence of automated and simplified scanning solutions "
    "limits the ability to continuously monitor and detect vulnerabilities in a "
    "timely manner, exposing systems to prolonged security risks "
    "(Tracy et al., 2007).",
    STATS)

# ------------------------------------------------- reference list
refs_cell = doc.tables[2].rows[7].cells[0]
ref_paras = [p for p in refs_cell.paragraphs if p.text.strip()]
existing = {norm(p.text) for p in ref_paras}

CVSS_REF = ("Forum of Incident Response and Security Teams. (2019). Common "
            "Vulnerability Scoring System v3.1: Specification document. "
            "https://www.first.org/cvss/v3.1/specification-document")

if not any("v3.1: Specification document" in t for t in existing):
    template = ref_paras[-1]._p
    new_p = copy.deepcopy(template)
    template.addnext(new_p)
    set_paragraph_text(Paragraph(new_p, refs_cell), CVSS_REF)
    ref_paras.append(Paragraph(new_p, refs_cell))
    STATS.append("added reference: FIRST CVSS v3.1")

# sort the reference paragraphs alphabetically (APA), keeping formatting
keys = []
for p in ref_paras:
    key = norm(p.text).lstrip("\u200c\u200b\u200e\u200f \t").lower()
    keys.append(key)
new_order = [p for _, p in sorted(zip(keys, ref_paras), key=lambda t: t[0])]
anchor = ref_paras[0]._p.getprevious()
parent = ref_paras[0]._p.getparent()
els = [p._p for p in new_order]
for el in els:
    parent.remove(el)
insert_at = list(parent).index(anchor) + 1
for offset, el in enumerate(els):
    parent.insert(insert_at + offset, el)
STATS.append("reference list re-sorted alphabetically")

doc.save(SRC)
print()
for s in STATS:
    print(" -", s)
print()
print("edits applied:", len([s for s in STATS if not s.startswith("!!")]))
