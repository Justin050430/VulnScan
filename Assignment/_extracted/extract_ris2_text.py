"""Utility: dump RIS Form 2 (proposal) to plain text for review/diffing."""
from docx import Document

SRC = r"c:\FYP\VulnScan\Assignment\OoiJianXin_RIS_Form2.docx"
OUT = r"c:\FYP\VulnScan\Assignment\_extracted\ris2_full.txt"

doc = Document(SRC)
lines = []
for ti, table in enumerate(doc.tables):
    lines.append(f"===== TABLE {ti} ({len(table.rows)}x{len(table.columns)}) =====")
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            texts = [p.text.strip() for p in cell.paragraphs if p.text.strip()]
            if texts:
                lines.append(f"--- T{ti} R{ri} C{ci} ---")
                lines.extend(texts)
            for nt in cell.tables:
                lines.append(f"--- T{ti} R{ri} C{ci} nested ({len(nt.rows)}x{len(nt.columns)}) ---")
                for nrow in nt.rows:
                    lines.append("  | " + " | ".join(
                        c.text.strip().replace("\n", " ") for c in nrow.cells))
    lines.append("")

with open(OUT, "w", encoding="utf-8") as fh:
    fh.write("\n".join(lines))
print(f"wrote {OUT} ({len(lines)} lines)")
