# -*- coding: utf-8 -*-
"""
Replace the auth ACTIVITY figure (currently Figure 4.9) with the new
Figure 4.9 User Authentication Sequence Diagram, and update caption + purpose text.
"""
import shutil, docx
from PIL import Image
import io

SRC = r"c:\FYP\VulnScan\Assignment\FYP.docx"
PNG = r"c:\FYP\VulnScan\docs\design\Fig4.9_Auth_Sequence_Diagram.png"
shutil.copyfile(SRC, r"c:\FYP\VulnScan\Assignment\FYP_backup_before_fig49.docx")

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"

OLD_CAP = "Figure 4.9: User Authentication Activity Diagram"
NEW_CAP_LABEL = "Figure 4.9:"
NEW_CAP_REST = " User Authentication Sequence Diagram"

NEW_PURPOSE = ("This sequence diagram illustrates the user authentication process of VulnScan. It shows how a "
               "user submits their email address and password, how the request is validated by the Flask API and "
               "the authentication service, and how the service verifies the stored password hash and account "
               "status against the database before issuing a JSON Web Token (JWT). The opt fragment shows the "
               "alternative Google single sign-on flow, in which the verified email address returned by Google is "
               "matched against an existing user account before the same JWT is issued. Successful authentication "
               "redirects the user to the dashboard appropriate to their role, while failed attempts result in an "
               "error response and are recorded in the security audit log.")

d = docx.Document(SRC)
paras = d.paragraphs
cap = next(p for p in paras if p.text.strip().startswith("Figure 4.9: User Authentication Activity Diagram"))

# The old auth activity image no longer exists in the doc -> insert a new image paragraph before the caption
tpl = next(p for p in paras if p.style.name == "Normal" and len(p.text.strip()) > 150 and not p._p.findall('.//{%s}blip' % NS_A))
from docx.text.paragraph import Paragraph
from docx.shared import Inches
el = __import__("copy").deepcopy(tpl._p)
cap._p.addprevious(el)
np = Paragraph(el, tpl._parent)
for r in list(np.runs):
    r._element.getparent().remove(r._element)
np.add_run().add_picture(PNG, width=Inches(6.0))

# caption: bold label + normal name, centered
for r in list(cap.runs):
    r._element.getparent().remove(r._element)
r1 = cap.add_run(NEW_CAP_LABEL); r1.bold = True
r2 = cap.add_run(NEW_CAP_REST); r2.bold = False
from docx.enum.text import WD_ALIGN_PARAGRAPH
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

# purpose text
idx = next(i for i, p in enumerate(paras) if p._element is cap._element)
for j in range(idx + 1, min(idx + 4, len(paras))):
    p = paras[j]
    if "workflow illustrates the authentication process" in p.text or "login credentials are validated" in p.text:
        if p.runs:
            p.runs[0].text = NEW_PURPOSE
            for r in p.runs[1:]:
                r._element.getparent().remove(r._element)
        else:
            p.add_run(NEW_PURPOSE)
        print("purpose updated at", j)
        break

d.save(SRC)
print("saved", SRC)
