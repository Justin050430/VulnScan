# -*- coding: utf-8 -*-
"""Replace Figure 4.12 audit activity image with the new vertical-swimlane UML activity diagram."""
import shutil, docx, io
from copy import deepcopy
from PIL import Image
from docx.text.paragraph import Paragraph
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = r"c:\FYP\VulnScan\Assignment\FYP.docx"
DST = r"c:\FYP\VulnScan\Assignment\FYP_updated.docx"
PNG = r"c:\FYP\VulnScan\docs\design\Fig4.12_Audit_Activity_Diagram.png"
shutil.copyfile(SRC, r"c:\FYP\VulnScan\Assignment\FYP_backup_before_fig412.docx")

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NEW_PURPOSE = ("This activity diagram illustrates how VulnScan records security audit logs using two swimlanes. "
               "When a significant action occurs, such as authentication, scan creation, or an administrative "
               "operation, the Business Logic layer captures the event context, decides whether the event is "
               "significant, and composes an audit record containing the category, action, description, responsible "
               "user, and event time. Significant records are inserted into the SecurityAuditLog entity in the "
               "database, where administrators can subsequently review them to support monitoring and forensic "
               "analysis. Non-significant events are discarded without a record.")

d = docx.Document(SRC)
paras = d.paragraphs

def has_blip(p):
    return len(p._p.findall('.//{%s}blip' % NS_A)) > 0

cap = next(p for p in paras if p.text.strip().startswith("Figure 4.12: Security Audit Logging Activity Diagram"))
idx = next(i for i, p in enumerate(paras) if p._element is cap._element)

img = cap if has_blip(cap) else None
if img is None:
    for j in range(idx - 1, max(-1, idx - 6), -1):
        if has_blip(paras[j]):
            img = paras[j]
            break

with Image.open(PNG) as im:
    aspect = im.width / im.height
W_EMU = int(5.8 * 914400)
H_EMU = int(W_EMU / aspect)

if img is not None:
    blip = img._p.findall('.//{%s}blip' % NS_A)[0]
    embed = blip.get('{%s}embed' % NS_R)
    d.part.rels[embed].target_part._blob = open(PNG, "rb").read()
    for ext in img._p.iter('{%s}extent' % NS_WP):
        ext.set('cx', str(W_EMU)); ext.set('cy', str(H_EMU))
    for aext in blip.iter('{%s}ext' % NS_A):
        aext.set('cx', str(W_EMU)); aext.set('cy', str(H_EMU))
    print("replaced existing image")
else:
    tpl = next(p for p in paras if p.style.name == "Normal" and len(p.text.strip()) > 150 and not has_blip(p))
    el = deepcopy(tpl._p)
    cap._p.addprevious(el)
    np = Paragraph(el, tpl._parent)
    for r in list(np.runs):
        r._element.getparent().remove(r._element)
    np.add_run().add_picture(PNG, width=Inches(5.8))
    print("inserted new image paragraph")

# caption (bold label + regular name) - normalise
full = cap.text.strip()
i_colon = full.find(":")
label = full[:i_colon + 1]; rest = full[i_colon + 1:]
for r in list(cap.runs):
    r._element.getparent().remove(r._element)
r1 = cap.add_run(label); r1.bold = True
r2 = cap.add_run(rest); r2.bold = False
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

# purpose text
idx = next(i for i, p in enumerate(paras) if p._element is cap._element)
for j in range(idx + 1, min(idx + 4, len(paras))):
    p = paras[j]
    if "audit logs" in p.text and ("captured" in p.text or "records and manages" in p.text or "stored in the SecurityAuditLog" in p.text):
        if p.runs:
            p.runs[0].text = NEW_PURPOSE
            for r in p.runs[1:]:
                r._element.getparent().remove(r._element)
        else:
            p.add_run(NEW_PURPOSE)
        print("purpose updated at", j)
        break

try:
    d.save(SRC)
    print("saved", SRC)
except PermissionError:
    d.save(DST)
    print("SRC locked; saved", DST)
