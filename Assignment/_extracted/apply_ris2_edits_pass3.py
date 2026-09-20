"""Third pass: keep the technology/hosting wording truthful whether or not the
Azure deployment has been performed (the report describes it as the target
deployment)."""
from docx import Document

SRC = r"c:\FYP\VulnScan\Assignment\OoiJianXin_RIS_Form2.docx"
doc = Document(SRC)
tech = doc.tables[1].rows[7].cells[0].tables[0]


def set_cell(cell, new_text):
    p = cell.paragraphs[0]
    runs = p.runs
    runs[0].text = new_text
    for r in runs[1:]:
        r.text = ""


set_cell(tech.rows[3].cells[1], "Microsoft SQL Server")
set_cell(tech.rows[8].cells[2],
         "Planned hosting platform for the web application and its managed "
         "database, with automated deployment through a CI/CD pipeline")

doc.save(SRC)

tech = Document(SRC).tables[1].rows[7].cells[0].tables[0]
for row in tech.rows:
    print(" | ".join(c.text.strip() for c in row.cells))
