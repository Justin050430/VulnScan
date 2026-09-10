# -*- coding: utf-8 -*-
"""
Integrate Figure 4.8 (System Use Case Diagram) into FYP.docx:
 1) insert Heading 3 '4.4.1 System Use Case Model' + narrative + image + caption
    immediately before the current '4.4.1 User Authentication Process'
 2) renumber the following 4.4.x headings (+1)
 3) renumber following Figure captions (+1)
"""
import shutil, re, docx
from copy import deepcopy
from docx.text.paragraph import Paragraph
from docx.shared import Inches

SRC = r"c:\FYP\VulnScan\Assignment\FYP.docx"
PNG = r"c:\FYP\VulnScan\docs\design\Fig4.8_Use_Case_Diagram.png"
shutil.copyfile(SRC, r"c:\FYP\VulnScan\Assignment\FYP_backup_before_fig48.docx")

HEAD = "4.4.1 System Use Case Model"
CAP = "Figure 4.8: System Use Case Diagram of VulnScan"
NARR = ("The use case model describes the functional requirements of VulnScan from the perspective of its "
        "actors. As established in the requirements analysis, the system serves two roles: the Administrator and "
        "the Developer. The Administrator is responsible for management activities, including creating and "
        "maintaining user accounts, managing vulnerability definitions and scan profiles, and reviewing security "
        "audit logs. The Developer focuses on the vulnerability assessment workflow, including starting scans, "
        "viewing scan results, and generating or exporting scan reports. Both roles can sign in using their "
        "registered credentials; sign-in supports both local email and password authentication and Google single "
        "sign-on, and both roles can consult the vulnerability knowledge base. Figure 4.8 presents the use case "
        "diagram of VulnScan, which provides the functional basis for the detailed design in the remainder of this "
        "chapter.")

d = docx.Document(SRC)
paras = d.paragraphs

auth = next(p for p in paras if p.style.name == "Heading 3" and p.text.strip().startswith("4.4.1 User Authentication Process"))
nrm = next(p for p in paras if p.style.name == "Normal" and len(p.text.strip()) > 150 and not p._p.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip'))

def add_before(anchor, template_p, text=None, picture=None):
    el = deepcopy(template_p._p)
    anchor._p.addprevious(el)
    np = Paragraph(el, template_p._parent)
    if picture is None:
        if np.runs:
            np.runs[0].text = text
            for r in np.runs[1:]:
                r._element.getparent().remove(r._element)
        else:
            np.add_run(text)
    else:
        for r in list(np.runs):
            r._element.getparent().remove(r._element)
        np.add_run().add_picture(picture, width=Inches(6.0))
    return np

# insert in forward order (each addprevious(anchor) keeps forward order)
add_before(auth, auth, HEAD)          # heading 4.4.1 use case
add_before(auth, nrm, NARR)
add_before(auth, nrm, picture=PNG)
add_before(auth, nrm, CAP)

# renumber following 4.4.x headings (+1)
paras = d.paragraphs
i_head = next(i for i, p in enumerate(paras) if p.style.name == "Heading 3" and p.text.strip() == HEAD)
for i in range(i_head + 1, len(paras)):
    p = paras[i]
    if p.style.name == "Heading 3":
        m = re.match(r"^(4\.4\.)(\d+)( .*)$", p.text.strip())
        if m:
            newnum = int(m.group(2)) + 1
            newtext = "%s%d%s" % (m.group(1), newnum, m.group(3))
            if p.runs:
                p.runs[0].text = newtext
                for r in p.runs[1:]:
                    r._element.getparent().remove(r._element)
            else:
                p.add_run(newtext)

# renumber following Figure captions (+1)
i_cap = next(i for i, p in enumerate(paras) if p.text.strip() == CAP)
for i in range(i_cap + 1, len(paras)):
    p = paras[i]
    m = re.match(r"^(Figure 4\.)(\d+)(:.*)$", p.text.strip())
    if m:
        newnum = int(m.group(2)) + 1
        newtext = "%s%d%s" % (m.group(1), newnum, m.group(3))
        # rebuild caption: bold label + regular rest (already split runs) -> adjust numbers in both runs
        # simplest: set first run label bold, rest normal
        runs = p.runs
        if runs:
            full = newtext
            idx = full.find(":")
            lab = full[:idx + 1]
            rest = full[idx + 1:]
            runs[0].text = lab
            runs[0].bold = True
            if len(runs) > 1:
                runs[1].text = rest
                runs[1].bold = False
                for r in runs[2:]:
                    r._element.getparent().remove(r._element)
            else:
                r2 = p.add_run(rest)
                r2.bold = False

d.save(SRC)
print("saved", SRC)
