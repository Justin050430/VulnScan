# -*- coding: utf-8 -*-
"""
Rebuild the References section of FYP.docx as a proper APA author-year list,
and normalise the in-text Scambray citations to the verified 2006 source.
"""
import docx
from copy import deepcopy
from docx.text.paragraph import Paragraph

PATH = r"c:\FYP\VulnScan\Assignment\FYP.docx"
doc = docx.Document(PATH)
paras = doc.paragraphs

def set_para_text(p, text):
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        p.add_run(text)

log = []
# ---------------------------------------------------------- in-text years
for p in paras:
    if "Scambray et al., 2011" in p.text:
        set_para_text(p, p.text.replace("Scambray et al., 2011", "Scambray et al., 2006"))
        log.append("normalised 2.1.2 Scambray year -> 2006")
    if "Scambray et al., 2021" in p.text:
        set_para_text(p, p.text.replace("Scambray et al., 2021", "Scambray et al., 2006"))
        log.append("normalised 1.1.3 Scambray year -> 2006")

# ---------------------------------------------------------- locate headings
ref_heading = None
app_heading = None
for p in paras:
    t = p.text.strip()
    if ref_heading is None and p.style.name == "Heading 1" and t == "References":
        ref_heading = p
    if app_heading is None and t == "Appendices":
        app_heading = p
print("ref heading:", ref_heading is not None, "| appendices:", app_heading is not None)

# delete everything between References heading and Appendices heading
if ref_heading is not None and app_heading is not None:
    to_delete = []
    node = ref_heading._p.getnext()
    while node is not None and node is not app_heading._p:
        nxt = node.getnext()
        to_delete.append(node)
        node = nxt
    for el in to_delete:
        el.getparent().remove(el)
    log.append(f"deleted {len(to_delete)} paragraphs in old References block")

# ---------------------------------------------------------- new APA entries
REFERENCES = [
    "Beck, K., Beedle, M., van Bennekum, A., Cockburn, A., Cunningham, W., Fowler, M., Grenning, J., "
    "Highsmith, J., Hunt, A., Jeffries, R., Kern, J., Marick, B., Martin, R. C., Mellor, S., Schwaber, K., "
    "Sutherland, J., & Thomas, D. (2001). Manifesto for Agile software development. "
    "https://agilemanifesto.org/",
    "Ferraiolo, D. F., Sandhu, R., Gavrila, S., Kuhn, D. R., & Chandramouli, R. (2001). Proposed NIST standard "
    "for role-based access control. ACM Transactions on Information and System Security, 4(3), 224\u2013274. "
    "https://doi.org/10.1145/501978.501980",
    "Forum of Incident Response and Security Teams. (2023). Common Vulnerability Scoring System version 4.0. "
    "Retrieved 4 September 2026, from https://www.first.org/cvss/",
    "Invicti Security. (2025). Acunetix web application security scanner. Retrieved 4 September 2026, "
    "from https://www.acunetix.com/",
    "Mell, P., Scarfone, K., & Romanosky, S. (2007). A complete guide to the Common Vulnerability Scoring "
    "System version 2.0. Forum of Incident Response and Security Teams. https://www.first.org/cvss/v2/guide",
    "National Institute of Standards and Technology. (2020). Security and privacy controls for information "
    "systems and organizations (NIST Special Publication 800-53, Revision 5). "
    "https://doi.org/10.6028/NIST.SP.800-53r5",
    "National Institute of Standards and Technology. (2025). Role-based access control. "
    "Retrieved 4 September 2026, from https://csrc.nist.gov/projects/role-based-access-control",
    "Open Worldwide Application Security Project. (2021). OWASP Top 10:2021 \u2013 The ten most critical web "
    "application security risks. Retrieved 4 September 2026, from https://owasp.org/Top10/",
    "Open Worldwide Application Security Project. (2025). OWASP Zed Attack Proxy (ZAP). "
    "Retrieved 4 September 2026, from https://www.zaproxy.org/",
    "PortSwigger. (2025). Burp Suite documentation. Retrieved 4 September 2026, "
    "from https://portswigger.net/burp/documentation",
    "Pressman, R. S., & Maxim, B. R. (2020). Software engineering: A practitioner\u2019s approach (9th ed.). "
    "McGraw-Hill Education.",
    "Scambray, J., Shema, M., & Sima, C. (2006). Hacking exposed web applications (2nd ed.). "
    "McGraw-Hill/Osborne.",
    "Sommerville, I. (2016). Software engineering (10th ed.). Pearson Education.",
]

# clone formatting from a normal body paragraph
template = None
for p in paras:
    if p.style.name == "normal" and p.text.strip().startswith("This chapter presented the overall system design"):
        template = p
        break
if template is None:
    template = ref_heading  # fallback

anchor = ref_heading
for text in REFERENCES:
    new_el = deepcopy(template._p)
    anchor._p.addnext(new_el)
    np = Paragraph(new_el, template._parent)
    set_para_text(np, text)
    anchor = np

log.append(f"inserted {len(REFERENCES)} APA references")
doc.save(PATH)
print("\n".join(log))
print("Saved:", PATH)
