# -*- coding: utf-8 -*-
"""
Insert 4.4.4 Vulnerability Detection Modules (with Figure 4.11 SQLi flowchart)
before the audit section, then renumber following 4.4.x headings (+1) and Figure captions (+1).
"""
import shutil, re, docx
from copy import deepcopy
from docx.text.paragraph import Paragraph
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = r"c:\FYP\VulnScan\Assignment\FYP.docx"
PNG = r"c:\FYP\VulnScan\docs\design\Fig4.11_SQLi_Flowchart.png"
shutil.copyfile(SRC, r"c:\FYP\VulnScan\Assignment\FYP_backup_before_fig411.docx")

HEAD = "4.4.4 Vulnerability Detection Modules"
CAP_LABEL = "Figure 4.11:"
CAP_REST = " SQL Injection Detection Flowchart"
NARR = ("VulnScan supports three vulnerability detection modules in its initial version: SQL Injection (SQLi), "
        "Cross-Site Scripting (XSS), and HTTP Security Header analysis. These modules are executed against the "
        "pages discovered by the Web Crawler according to the modules enabled in the selected scan profile, and "
        "they are represented by the Vulnerability and ProfileModule entities in the database design. The SQL "
        "Injection detection module tests the input parameters of each page by sending crafted payloads and "
        "analysing the responses for SQL errors or anomalous behaviour; when a vulnerability is indicated, a "
        "Finding is recorded together with its severity, risk score, and supporting evidence. Figure 4.11 presents "
        "the flowchart of the SQL Injection detection algorithm. The XSS module applies a similar payload and "
        "response analysis to detect whether untrusted input is reflected and executed in the returned page, while "
        "the HTTP Security Header module inspects the response headers of each page against recommended security "
        "header settings. The detection modules are implemented as independent components, allowing additional "
        "modules to be added in the future without modifying the core scanner architecture.")

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
d = docx.Document(SRC)
paras = d.paragraphs

audit = next(p for p in paras if p.style.name == "Heading 3" and p.text.strip().startswith("4.4.4 Security Audit Logging Process"))
nrm = next(p for p in paras if p.style.name == "Normal" and len(p.text.strip()) > 150 and not p._p.findall('.//{%s}blip' % NS_A))

def add_before(anchor, template_p, text=None, picture=None):
    el = deepcopy(template_p._p)
    anchor._p.addprevious(el)
    np = Paragraph(el, template_p._parent)
    if picture is None:
        if np.runs:
            np.runs[0].text = text
            for r in np.runs[1:]:
                r._element.getparent().remove(r._element)
        else:
            np.add_run(text)
    else:
        for r in list(np.runs):
            r._element.getparent().remove(r._element)
        np.add_run().add_picture(picture, width=Inches(5.6))
    return np

add_before(audit, audit, HEAD)
add_before(audit, nrm, NARR)
add_before(audit, nrm, picture=PNG)
cap = add_before(audit, nrm, "")
# caption: bold label + normal name, centred
for r in list(cap.runs):
    r._element.getparent().remove(r._element)
r1 = cap.add_run(CAP_LABEL); r1.bold = True
r2 = cap.add_run(CAP_REST); r2.bold = False
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

# renumber subsequent 4.4.x headings (+1)
paras = d.paragraphs
i_head = next(i for i, p in enumerate(paras) if p.style.name == "Heading 3" and p.text.strip() == HEAD)
for i in range(i_head + 1, len(paras)):
    p = paras[i]
    if p.style.name == "Heading 3":
        m = re.match(r"^(4\.4\.)(\d+)( .*)$", p.text.strip())
        if m and int(m.group(2)) >= 4:
            newnum = int(m.group(2)) + 1
            newtext = "%s%d%s" % (m.group(1), newnum, m.group(3))
            if p.runs:
                p.runs[0].text = newtext
                for r in p.runs[1:]:
                    r._element.getparent().remove(r._element)
            else:
                p.add_run(newtext)

# renumber subsequent Figure captions (+1)
i_cap = next(i for i, p in enumerate(paras) if p.text.strip().startswith(CAP_LABEL))
for i in range(i_cap + 1, len(paras)):
    p = paras[i]
    m = re.match(r"^(Figure 4\.)(\d+)(:.*)$", p.text.strip())
    if m:
        newnum = int(m.group(2)) + 1
        full = "%s%d%s" % (m.group(1), newnum, m.group(3))
        idx = full.find(":")
        lab = full[:idx + 1]; rest = full[idx + 1:]
        runs = p.runs
        if runs:
            runs[0].text = lab; runs[0].bold = True
            if len(runs) > 1:
                runs[1].text = rest; runs[1].bold = False
                for r in runs[2:]:
                    r._element.getparent().remove(r._element)
            else:
                rr = p.add_run(rest); rr.bold = False

d.save(SRC)
print("saved", SRC)
