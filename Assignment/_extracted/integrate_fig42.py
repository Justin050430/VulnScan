# -*- coding: utf-8 -*-
"""
Integrate Figure 4.2 (Deployment Diagram) into FYP.docx:
 1) insert Heading 3 '4.1.2 Deployment Architecture' + narrative + image + caption
    just before '4.2 Database Design'
 2) renumber every later Figure caption by +1 (ERD->4.3, UI->4.4.., activities->4.8..)
 3) update the in-text ERD cross-reference 'Figure 4.2' -> 'Figure 4.3'
"""
import shutil, re, docx
from copy import deepcopy
from docx.text.paragraph import Paragraph
from docx.shared import Inches

SRC = r"c:\FYP\VulnScan\Assignment\FYP.docx"
PNG = r"c:\FYP\VulnScan\docs\design\Fig4.2_Deployment_Diagram.png"

shutil.copyfile(SRC, r"c:\FYP\VulnScan\Assignment\FYP_backup_before_fig42.docx")

d = docx.Document(SRC)
paras = d.paragraphs
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"

HEAD412 = "4.1.2 Deployment Architecture"
NARR = ("In addition to the logical layered view, the physical deployment of VulnScan must be considered. "
        "VulnScan is designed as a client-server web application. Users access the system through a web browser, "
        "which communicates with the Flask application server over HTTPS. The application server hosts the "
        "Presentation, API, and Business Logic layers together with the Scanner Engine, and connects to a Microsoft "
        "SQL Server database server over TCP through the SQLAlchemy object-relational mapper. When an assessment is "
        "performed, the Scanner Engine sends HTTP or HTTPS requests to the target web application being evaluated. "
        "When a user signs in through Google single sign-on, the application server communicates with the Google "
        "OAuth 2.0 authorisation service over HTTPS. Figure 4.2 presents the deployment architecture of VulnScan.")
CAPTION = "Figure 4.2: Deployment Diagram of VulnScan"

# anchor: Heading 2 '4.2 Database Design'
anchor = next(p for p in paras if p.style.name == "Heading 2" and p.text.strip().startswith("4.2 Database Design"))
# templates
h3t = next(p for p in paras if p.style.name == "Heading 3" and p.text.strip().startswith("4.4.1 User Authentication"))
nrm = next(p for p in paras if p.style.name == "Normal" and len(p.text.strip()) > 120 and not p._p.findall('.//{%s}blip' % NS_A))


def clone(style_p):
    el = deepcopy(style_p._p)
    anchor._p.addprevious(el)
    return Paragraph(el, style_p._parent)

# insert in forward order before anchor
ph = clone(h3t)
if ph.runs:
    ph.runs[0].text = HEAD412
    for r in ph.runs[1:]:
        r._element.getparent().remove(r._element)
else:
    ph.add_run(HEAD412)

pn = clone(nrm)
if pn.runs:
    pn.runs[0].text = NARR
    for r in pn.runs[1:]:
        r._element.getparent().remove(r._element)
else:
    pn.add_run(NARR)

pi = clone(nrm)   # empty image paragraph
if pi.runs:
    for r in list(pi.runs):
        r._element.getparent().remove(r._element)
pi.add_run().add_picture(PNG, width=Inches(6.0))

pc = clone(nrm)
if pc.runs:
    pc.runs[0].text = CAPTION
    for r in pc.runs[1:]:
        r._element.getparent().remove(r._element)
else:
    pc.add_run(CAPTION)

# renumber every later Figure caption by +1 (skip our new caption which is at position 4.1.2)
paras = d.paragraphs
i_cap = next(i for i, p in enumerate(paras) if p.text.strip() == CAPTION)
n_bump = 0
for i in range(i_cap + 1, len(paras)):
    p = paras[i]
    m = re.match(r"^Figure 4\.(\d+):", p.text.strip())
    if m:
        old = int(m.group(1))
        newn = old + 1
        setattr(p, "_newtext", re.sub(r"^Figure 4\.\d+:", "Figure 4.%d:" % newn, p.text))
        if p.runs:
            p.runs[0].text = re.sub(r"^Figure 4\.\d+:", "Figure 4.%d:" % newn, p.text)
            for r in p.runs[1:]:
                r._element.getparent().remove(r._element)
        else:
            p.add_run(re.sub(r"^Figure 4\.\d+:", "Figure 4.%d:" % newn, p.text))
        n_bump += 1
print("renumbered captions:", n_bump)

# update ERD in-text cross reference
for p in d.paragraphs:
    if "As illustrated in Figure 4.2, the VulnScan database" in p.text:
        newt = p.text.replace("As illustrated in Figure 4.2, the VulnScan database",
                              "As illustrated in Figure 4.3, the VulnScan database")
        if p.runs:
            p.runs[0].text = newt
            for r in p.runs[1:]:
                r._element.getparent().remove(r._element)
        else:
            p.add_run(newt)
        print("updated ERD cross-reference")
        break

d.save(SRC)
print("saved", SRC)
