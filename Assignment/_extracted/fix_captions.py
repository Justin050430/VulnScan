# -*- coding: utf-8 -*-
"""
Normalise figure/table captions to the report convention:
  - centred
  - label token 'Figure X.Y:' (or 'Table X.Y:') BOLD, colon included
  - description text regular weight
Applies to all Figure/Table caption paragraphs.
"""
import docx, re, shutil
from copy import deepcopy
from docx.enum.text import WD_ALIGN_PARAGRAPH

P = r"c:\FYP\VulnScan\Assignment\FYP.docx"
shutil.copyfile(P, r"c:\FYP\VulnScan\Assignment\FYP_backup_before_captions.docx")

d = docx.Document(P)
pat = re.compile(r"^(Figure|Table) (\d+\.\d+):(.*)$", re.S)

changed = 0
for p in d.paragraphs:
    m = pat.match(p.text.strip())
    if not m:
        continue
    kind, num, rest = m.group(1), m.group(2), m.group(3)
    label = f"{kind} {num}:"
    # rest: strip but keep a single leading space for readability
    rest_clean = " " + rest.strip() if rest.strip() else ""
    # capture rPr from first run if any
    rpr = None
    if p.runs:
        r0 = p.runs[0]
        if r0._r.rPr is not None:
            rpr = deepcopy(r0._r.rPr)
    # clear runs
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    r1 = p.add_run(label)
    r1.bold = True
    r2 = p.add_run(rest_clean)
    r2.bold = False
    if rpr is not None:
        for r in (r1, r2):
            if r._r.rPr is not None:
                r._r.remove(r._r.rPr)
            r._r.insert(0, deepcopy(rpr))
    # reapply bold toggles (bold may have been part of rpr; override)
    r1.bold = True
    r2.bold = False
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    changed += 1

d.save(P)
print("caption paragraphs normalised:", changed)
