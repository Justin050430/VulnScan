# -*- coding: utf-8 -*-
"""
VulnScan FYP.docx — Chapters 1-4 polish (batch A remaining items).
Items:
 A1 layer-name standardisation (match Fig 4.1: API Layer / Scanner Engine Layer / Data Layer)
 A2 objectives 1 & 4 distinct verbs (develop / implement)
 A3 3.1.3 heading singular
 A4 Table 3.3 + §3.3.1 prose include new FRs
 A5 new §4.4.4 API design + §4.4.5 design rationale
 A7 retitle §1.1.3 / §1.1.4 to match content
"""
import docx
from copy import deepcopy
from docx.text.paragraph import Paragraph
from docx.shared import Pt

PATH = r"c:\FYP\VulnScan\Assignment\FYP.docx"
doc = docx.Document(PATH)
paras = doc.paragraphs
log = []

def set_para_text(p, text):
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        p.add_run(text)

def replace_in(p, old, new):
    if old in p.text:
        set_para_text(p, p.text.replace(old, new))
        return True
    return False

# ---------------------------------------------------------------- A1: layer naming
GLOBAL_SUBS = [
    ("REST API Layer", "API Layer"),
    ("Database Layer", "Data Layer"),
    ("Scanner Engine, and Data Layer", "Scanner Engine Layer, and Data Layer"),
]
count = 0
for p in paras:
    for old, new in GLOBAL_SUBS:
        if replace_in(p, old, new):
            count += 1
log.append(f"A1: applied {count} layer-name substitutions")

# ---------------------------------------------------------------- A2: objectives verbs
def replace_para_starts_with(prefix, new_text):
    ms = [p for p in paras if p.text.strip().startswith(prefix)]
    if len(ms) != 1:
        log.append(f"!! A2 obj not unique: '{prefix[:40]}' -> {len(ms)}")
        return
    set_para_text(ms[0], new_text)
    log.append(f"A2: reworded objective '{prefix[:35]}...'")

replace_para_starts_with(
    "To design a web-based vulnerability assessment platform capable of automatically scanning web applications",
    "To develop a web-based vulnerability assessment platform capable of automatically scanning web applications "
    "for common security vulnerabilities based on selected OWASP web security practices.")
replace_para_starts_with(
    "To design a centralized management system that supports user authentication",
    "To implement a centralized management system that supports user authentication, role-based access control, "
    "configurable scan profiles, vulnerability management, and comprehensive security audit logging.")

# ---------------------------------------------------------------- A3: 3.1.3 heading singular
for p in paras:
    if p.style.name == "Heading 3" and p.text.strip() == "3.1.3 Design Phases":
        set_para_text(p, "3.1.3 Design Phase")
        log.append("A3: heading -> '3.1.3 Design Phase'")

# ---------------------------------------------------------------- A4: prioritisation table + prose
for ti, tbl in enumerate(doc.tables):
    hdr = [c.text.strip() for c in tbl.rows[0].cells]
    if hdr and hdr[0].lower() == "priority":
        for row in tbl.rows:
            key = row.cells[0].text.strip().lower()
            if key == "high":
                set_para_text(row.cells[1].paragraphs[0],
                    "User Authentication, Role-Based Access Control, User Management, Web Crawling, "
                    "Vulnerability Detection, Scan Result Management")
                log.append("A4: Table High row updated")
            elif key == "medium":
                set_para_text(row.cells[1].paragraphs[0],
                    "Scan Profile Management, Vulnerability Management, Single Sign-On (Google), "
                    "Scan Report Generation and Export, Security Audit Logging")
                log.append("A4: Table Medium row updated")

for p in paras:
    if p.text.strip().startswith("The highest priority requirements include user authentication"):
        set_para_text(p,
            "The highest priority requirements include user authentication, role-based access control, user "
            "management, vulnerability scanning, web crawling, scan result management, and vulnerability detection, "
            "as these modules form the core functionality of VulnScan. Supporting features such as scan profile "
            "management, vulnerability management, Google single sign-on, security audit logging, and scan report "
            "generation and export are also important but are implemented after the core scanning functions have "
            "been established. The vulnerability knowledge base, which enriches scan results with remediation "
            "guidance, is expanded continuously throughout development.")
        log.append("A4: §3.3.1 prose updated")

# ---------------------------------------------------------------- A7: retitle 1.1.3 / 1.1.4
for p in paras:
    t = p.text.strip()
    if p.style.name == "Heading 3" and t == "1.1.3 Web Vulnerabilities":
        set_para_text(p, "1.1.3 Vulnerability Assessment")
        log.append("A7: 1.1.3 heading -> 'Vulnerability Assessment'")
    elif p.style.name == "Heading 3" and t == "1.1.4 Vulnerability Assessment":
        set_para_text(p, "1.1.4 Web Vulnerability Standards (OWASP Top 10)")
        log.append("A7: 1.1.4 heading -> 'Web Vulnerability Standards (OWASP Top 10)'")

# ---------------------------------------------------------------- A5: API design + rationale
h45 = [p for p in paras if p.style.name == "Heading 2" and p.text.strip().startswith("4.5 Chapter Summary")]
h43 = [p for p in paras if p.style.name == "Heading 3" and p.text.strip().startswith("4.4.3 Security Audit")]
norm = [p for p in paras if p.style.name == "normal" and p.text.strip().startswith("The audit logging mechanism")]
if len(h45) == 1 and len(h43) == 1 and len(norm) == 1:
    anchor = h45[0]
    h3_tpl = h43[0]._p
    n_tpl = norm[0]._p

    def clone_p(style_el, text):
        el = deepcopy(style_el)
        np = Paragraph(el, anchor._parent)
        set_para_text(np, text)
        return el

    elements = []
    # (build in final order)
    elements.append(clone_p(h3_tpl, "4.4.4 Application Programming Interface (API) Design"))
    elements.append(clone_p(n_tpl,
        "VulnScan exposes its core services through a representational state transfer (REST) application "
        "programming interface (API). The Presentation Layer consumes this API, which is processed by the API "
        "Layer and returns responses in JavaScript Object Notation (JSON) format. Access to the API is protected "
        "using the JSON Web Tokens (JWT) described in Section 4.4.1, which clients include as bearer tokens in the "
        "Authorization header of each request. Permissions for each endpoint are enforced according to the role of "
        "the authenticated user, consistent with the role-based access control model presented in this chapter. "
        "Table 4.2 summarises the main endpoints exposed by the VulnScan API."))
    elements.append(clone_p(n_tpl, "Table 4.2: REST API Endpoints of VulnScan"))

    # build the API table
    rows_data = [
        ("Method", "Endpoint", "Description", "Access"),
        ("POST", "/api/auth/login", "Authenticates a registered user using email address and password and returns a JWT", "Public"),
        ("POST", "/api/auth/google", "Completes Google OAuth 2.0 sign-in and returns a JWT", "Public"),
        ("POST", "/api/auth/logout", "Invalidates the current user session", "Developer, Administrator"),
        ("GET", "/api/users", "Lists user accounts", "Administrator"),
        ("POST", "/api/users", "Creates a new user account and assigns a role", "Administrator"),
        ("PUT", "/api/users/{user_id}", "Updates user account details or role", "Administrator"),
        ("DELETE", "/api/users/{user_id}", "Disables or deletes a user account", "Administrator"),
        ("GET", "/api/vulnerabilities", "Retrieves vulnerability knowledge-base entries", "Developer, Administrator"),
        ("POST", "/api/vulnerabilities", "Creates a vulnerability definition", "Administrator"),
        ("PUT", "/api/vulnerabilities/{vulnerability_id}", "Updates a vulnerability definition", "Administrator"),
        ("GET", "/api/scan-profiles", "Lists the available scan profiles", "Developer, Administrator"),
        ("POST", "/api/scan-profiles", "Creates a scan profile", "Administrator"),
        ("PUT", "/api/scan-profiles/{profile_id}", "Updates a scan profile", "Administrator"),
        ("DELETE", "/api/scan-profiles/{profile_id}", "Deletes a scan profile", "Administrator"),
        ("POST", "/api/scans", "Starts a vulnerability scan for a target URL using a selected scan profile", "Developer"),
        ("GET", "/api/scans", "Lists the scans performed by the authenticated user", "Developer"),
        ("GET", "/api/scans/{scan_id}", "Retrieves scan details, discovered pages, and findings", "Developer, Administrator"),
        ("GET", "/api/scans/{scan_id}/report", "Generates a scan report and exports the scan results in PDF or CSV format", "Developer, Administrator"),
        ("GET", "/api/audit-logs", "Retrieves security audit log records", "Administrator"),
    ]
    api_tbl = doc.add_table(rows=len(rows_data), cols=4)
    # reuse an existing table style for borders
    try:
        style_name = doc.tables[3].style.name if doc.tables[3].style is not None else None
        if style_name:
            api_tbl.style = style_name
    except Exception:
        pass
    ref_rpr = None
    try:
        r0 = doc.tables[3].rows[1].cells[0].paragraphs[0].runs[0]
        ref_rpr = deepcopy(r0._r.rPr) if r0._r.rPr is not None else None
    except Exception:
        ref_rpr = None

    def set_cell(cell, text, bold=False):
        p0 = cell.paragraphs[0]
        for r in list(p0.runs):
            r._element.getparent().remove(r._element)
        run = p0.add_run(text)
        run.bold = bold
        if ref_rpr is not None:
            if run._r.rPr is not None:
                run._r.remove(run._r.rPr)
            run._r.insert(0, deepcopy(ref_rpr))
    for ri, row_vals in enumerate(rows_data):
        for ci, val in enumerate(row_vals):
            set_cell(api_tbl.rows[ri].cells[ci], val, bold=(ri == 0))
    elements.append(api_tbl._tbl)

    elements.append(clone_p(n_tpl,
        "The endpoints correspond to the functional requirements and database entities described earlier in this "
        "chapter. This resource-oriented design keeps the user interface and backend services loosely coupled and "
        "supports future enhancement, such as the addition of new vulnerability detection modules or report formats, "
        "without altering the overall API structure."))

    elements.append(clone_p(h3_tpl, "4.4.5 Design Rationale and Trade-offs"))
    elements.append(clone_p(n_tpl,
        "The design decisions for VulnScan balance technical feasibility, maintainability, and security. A layered "
        "architecture separates presentation, business logic, and data concerns, which simplifies development, "
        "testing, and maintenance; the main trade-off is the additional communication between layers, which is "
        "acceptable for the scale of this system. The relational database is normalised to Third Normal Form (3NF) "
        "to reduce redundancy and prevent update anomalies, at the cost of requiring joins when retrieving data; "
        "given the modest data volumes produced by vulnerability scans, this overhead is negligible. JWT-based "
        "stateless authentication was preferred over server-side session storage because it avoids maintaining "
        "session state between requests and is well suited to the client-server REST architecture. Flask was "
        "selected for its lightweight and modular structure, while Microsoft SQL Server was chosen for its "
        "robustness, transactional reliability, and integration with the SQLAlchemy object-relational mapper. "
        "Finally, the scanner engine is designed as an independent module with pluggable detection modules so that "
        "future vulnerability checks can be added without redesigning the core system."))

    for el in reversed(elements):
        anchor._p.addprevious(el)
    log.append("A5: inserted §4.4.4 (API design + Table 4.2) and §4.4.5 (design rationale)")
else:
    log.append(f"!! A5 anchors not found (h45={len(h45)} h43={len(h43)} norm={len(norm)})")

doc.save(PATH)
print("\n".join(log))
print("Saved:", PATH)
