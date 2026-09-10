# -*- coding: utf-8 -*-
"""Replace Figure 4.10 scanning activity image with the new UML activity diagram."""
import shutil, docx, io
from copy import deepcopy
from PIL import Image
from docx.text.paragraph import Paragraph
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = r"c:\FYP\VulnScan\Assignment\FYP.docx"
PNG = r"c:\FYP\VulnScan\docs\design\Fig4.10_Scan_Activity_Diagram.png"
shutil.copyfile(SRC, r"c:\FYP\VulnScan\Assignment\FYP_backup_before_fig410.docx")

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"

d = docx.Document(SRC)
paras = d.paragraphs

def has_blip(p):
    return len(p._p.findall('.//{%s}blip' % NS_A)) > 0

cap = next(p for p in paras if p.text.strip().startswith("Figure 4.10: Vulnerability Scanning Activity Diagram"))
idx = next(i for i, p in enumerate(paras) if p._element is cap._element)

img = cap if has_blip(cap) else None
if img is None:
    for j in range(idx - 1, max(-1, idx - 6), -1):
        if has_blip(paras[j]):
            img = paras[j]
            break

with Image.open(PNG) as im:
    aspect = im.width / im.height
W_EMU = int(5.4 * 914400)
H_EMU = int(W_EMU / aspect)

if img is not None:
    blip = img._p.findall('.//{%s}blip' % NS_A)[0]
    embed = blip.get('{%s}embed' % NS_R)
    d.part.rels[embed].target_part._blob = open(PNG, "rb").read()
    for ext in img._p.iter('{%s}extent' % NS_WP):
        ext.set('cx', str(W_EMU)); ext.set('cy', str(H_EMU))
    for aext in blip.iter('{%s}ext' % NS_A):
        aext.set('cx', str(W_EMU)); aext.set('cy', str(H_EMU))
    print("replaced existing image")
else:
    tpl = next(p for p in paras if p.style.name == "Normal" and len(p.text.strip()) > 150 and not has_blip(p))
    el = deepcopy(tpl._p)
    cap._p.addprevious(el)
    np = Paragraph(el, tpl._parent)
    for r in list(np.runs):
        r._element.getparent().remove(r._element)
    np.add_run().add_picture(PNG, width=Inches(5.4))
    print("inserted new image paragraph")

# normalise caption (bold label + regular name)
full = cap.text.strip()
i_colon = full.find(":")
label = full[:i_colon + 1]
rest = full[i_colon + 1:]
for r in list(cap.runs):
    r._element.getparent().remove(r._element)
r1 = cap.add_run(label); r1.bold = True
r2 = cap.add_run(rest); r2.bold = False
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

d.save(SRC)
print("saved", SRC)
