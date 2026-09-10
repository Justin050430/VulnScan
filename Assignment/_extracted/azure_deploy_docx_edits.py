# -*- coding: utf-8 -*-
"""
Azure deployment consistency edits for FYP.docx (Chapters 3 & 4).
Operates on Assignment/FYP.docx (backup FYP_backup_before_azure.docx exists).
  1) Chapter 3 - Table 3.2 NFR: amend NFR-07 / NFR-08 and append NFR-09 (Deployability)
  2) Chapter 4 - 4.1.2 Deployment Architecture narrative -> Microsoft Azure + GitHub Actions
  3) Chapter 4 - 4.2 Database Design paragraph -> Azure SQL Database mention
  4) Swap the Figure 4.2 image bytes with the updated Fig4.2_Deployment_Diagram.png
Run from the repo venv.
"""
import docx, io
from copy import deepcopy
from docx.text.paragraph import Paragraph

PATH = r"c:\FYP\VulnScan\Assignment\FYP.docx"
PNG  = r"c:\FYP\VulnScan\docs\design\Fig4.2_Deployment_Diagram.png"
NS_A  = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_R  = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

doc = docx.Document(PATH)
report = []
def log(m): report.append(m)

# ---------------------------------------------------------------- helpers
def set_para_text(p, text):
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        p.add_run(text)

def find_paras(sub):
    return [p for p in doc.paragraphs if sub in p.text]

def replace_para(sub, new_text):
    ms = find_paras(sub)
    if len(ms) != 1:
        log(f"!! replace '{sub[:50]}' matched {len(ms)} -> SKIPPED")
        return False
    set_para_text(ms[0], new_text)
    log(f"OK  paragraph updated [{sub[:45]}...]")
    return True

def set_cell(cell, text):
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    set_para_text(cell.paragraphs[0], text)

# ================================================================ 1) NFR table
tbl = doc.tables[4]   # Table 3.2 Non-Functional Requirements
rows = tbl.rows
NFR_UPD = {
    "NFR-07": "The system shall operate on modern web browsers over HTTPS and support integration "
              "with Microsoft SQL Server, provisioned as Azure SQL Database, through the Flask backend.",
    "NFR-08": "The system shall remain accessible through a public HTTPS endpoint whenever required "
              "during normal operation and recover gracefully from unexpected failures where possible.",
}
for r in rows:
    rid = r.cells[0].text.strip()
    if rid in NFR_UPD:
        set_cell(r.cells[2], NFR_UPD[rid])
        log(f"OK  {rid} description updated")

# append NFR-09 (clone last data row to preserve formatting) - guarded so re-runs are safe
if any(r.cells[0].text.strip() == "NFR-09" for r in tbl.rows):
    log(".. NFR-09 already present - append skipped")
else:
    last_row_el = rows[-1]._tr
    new_tr = deepcopy(last_row_el)
    last_row_el.addnext(new_tr)
    tbl_row = None
    for r2 in tbl.rows:
        if r2._tr is new_tr:
            tbl_row = r2
            break
    set_cell(tbl_row.cells[0], "NFR-09")
    set_cell(tbl_row.cells[1], "Deployability")
    set_cell(tbl_row.cells[2],
            "The system shall be deployable as a cloud-hosted web application on Microsoft Azure and shall "
            "be updated automatically through a continuous integration and continuous deployment (CI/CD) "
            "pipeline whenever source code changes are pushed to the GitHub repository.")
    log("OK  appended NFR-09 (Deployability)")

# ---------------------------------------------------------------- 2) 4.1.2 narrative
NARR412 = (
    "In addition to the logical layered view, the physical deployment of VulnScan must be considered. "
    "VulnScan is designed as a client-server web application and is targeted for deployment on the "
    "Microsoft Azure cloud platform. Users access the system through a web browser over HTTPS using the "
    "public URL of the deployed application. The Flask web application, which serves both the frontend "
    "user interface and the backend REST API together with the Scanner Engine, is hosted as an Azure "
    "App Service. The application connects to Microsoft SQL Server through the SQLAlchemy "
    "object-relational mapper; in the target deployment, the database is provisioned as Azure SQL "
    "Database, a managed cloud service based on the Microsoft SQL Server engine. When an assessment is "
    "performed, the Scanner Engine sends HTTP or HTTPS requests to the target web application being "
    "evaluated. When a user signs in through Google single sign-on, the application communicates with "
    "the Google OAuth 2.0 authorisation service over HTTPS. To support continuous delivery, the source "
    "code is maintained in a GitHub repository, and a GitHub Actions continuous integration and "
    "continuous deployment (CI/CD) pipeline automatically redeploys the application to Azure App "
    "Service whenever changes are pushed to the repository. Figure 4.2 presents the deployment "
    "architecture of VulnScan.")
replace_para("the physical deployment of VulnScan must be considered", NARR412)

# ---------------------------------------------------------------- 3) 4.2 DB paragraph
DB412 = (
    "The database is designed to provide a structured and reliable repository for managing user "
    "information, scan configurations, vulnerability definitions, scan results, and system activities. "
    "Microsoft SQL Server was selected as the database management system because of its robustness, "
    "transactional reliability, advanced security features, and seamless integration with SQLAlchemy. "
    "These capabilities make it suitable for managing the relational data required by VulnScan while "
    "supporting future scalability and maintenance. During deployment, the database is provisioned as "
    "Azure SQL Database, a fully managed cloud service built on the Microsoft SQL Server engine that "
    "provides automated backups, high availability, and elastic scalability.")
replace_para("Microsoft SQL Server was selected as the database management system because of its "
             "robustness", DB412)

# ---------------------------------------------------------------- 4) swap Figure 4.2 image
def has_blip(p):
    return len(p._p.findall('.//{%s}blip' % NS_A)) > 0

cap = next((p for p in doc.paragraphs
            if p.text.strip().startswith("Figure 4.2: Deployment Diagram of VulnScan")), None)
assert cap is not None, "Figure 4.2 caption not found"
idx = next(i for i, p in enumerate(doc.paragraphs) if p._element is cap._element)
img = None
for j in range(idx - 1, max(-1, idx - 8), -1):
    if has_blip(doc.paragraphs[j]):
        img = doc.paragraphs[j]
        break
if img is not None:
    blip = img._p.findall('.//{%s}blip' % NS_A)[0]
    embed = blip.get('{%s}embed' % NS_R)
    doc.part.rels[embed].target_part._blob = open(PNG, "rb").read()
    log("OK  replaced Figure 4.2 image bytes with Azure diagram")
else:
    log("!! Figure 4.2 image paragraph not found - bytes NOT swapped")

# scan for any leftover physical-server wording worth knowing about
for p in doc.paragraphs:
    t = p.text
    if ("application server" in t.lower() and "Figure 4.2" not in t) or "Database Server" in t:
        log(f"NOTE other 'server' wording: [{p.text[:90]}...]")

doc.save(PATH)
print("\n".join(report))
print("\nSaved:", PATH)
